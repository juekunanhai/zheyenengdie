"""Export the current SPEC into both Word copies, retaining the historical appendix."""
from copy import deepcopy
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def plain(text):
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')


def inline(paragraph, text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(plain(part))
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def add_table(doc, lines, expected):
    rows = [[plain(cell.strip()) for cell in line.strip().strip('|').split('|')]
            for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-{3,}:?', c) for c in row)]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    available = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    widths = [available / len(rows[0])] * len(rows[0])
    if len(rows[0]) == 3:
        widths = [available * ratio for ratio in (0.25, 0.50, 0.25)]
    if len(rows[0]) == 2:
        widths = [available * ratio for ratio in (0.25, 0.75)]
    for column, width in zip(table.columns, widths):
        column.width = int(width)
    for i, values in enumerate(rows):
        row = table.add_row()
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = int(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            expected.append(value)
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    item.set(qn('w:' + key), val)
                borders.append(item)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), '80')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7EEF4')
                props.append(shade)
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)


def render_body(doc, markdown, expected):
    lines = markdown.splitlines()
    i = 0
    code = False
    while i < len(lines):
        line = lines[i]
        i += 1
        if line.startswith('```'):
            code = not code
            continue
        if code:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
            expected.append(line)
            continue
        if not line.strip() or line.strip() == '---':
            continue
        if line.startswith('# '):
            continue
        if line.startswith('|'):
            table_lines = [line]
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines, expected)
            continue
        heading = re.match(r'^(#{2,4})\s+(.+)', line)
        if heading:
            text = plain(heading[2])
            p = doc.add_paragraph(text, style='Heading ' + str(len(heading[1]) - 1))
            expected.append(text)
            continue
        bullet = re.match(r'^(\s*)- (.+)', line)
        if bullet:
            p = doc.add_paragraph(style='List Bullet')
            if bullet[1]:
                p.paragraph_format.left_indent = Inches(0.22 + 0.12 * len(bullet[1]))
            line = bullet[2]
        else:
            p = doc.add_paragraph()
            line = line.removeprefix('> ')
            if line.rstrip().endswith(('：', ':')):
                p.paragraph_format.keep_with_next = True
        inline(p, line)
        expected.append(plain(line))


def export(spec_path, word_path, second_path):
    doc = Document(word_path)
    body = doc._element.body
    appendix, retain = [], False
    for element in list(body):
        text = ''.join(element.itertext()) if element.tag == qn('w:p') else ''
        if '历史视觉参考 1' in text:
            retain = True
        if retain and element.tag != qn('w:sectPr'):
            appendix.append(deepcopy(element))
        if element.tag != qn('w:sectPr'):
            body.remove(element)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title = '这也能叠开发规格'
    subtitle = '跨平台 Cocos 游戏  微信首发'
    doc.add_paragraph(title, style='Title')
    doc.add_paragraph(subtitle)
    expected = [title, subtitle]
    render_body(doc, spec_path.read_text(), expected)
    doc.add_page_break()
    other_product = False
    for element in appendix:
        if element.tag == qn('w:p'):
            for node in element.findall('.//' + qn('w:t')):
                if node.text == '历史视觉参考 8':
                    other_product = True
                    node.text = '历史附图 8 其他产品图标'
                if node.text == '历史附图 8 其他产品图标':
                    other_product = True
                if node.text and node.text.startswith('仅用于理解视觉语言。'):
                    node.text = '仅用于理解历史视觉语言。当前布局以正文 R6 规则为准，图中旧功能、操作和数值不生效。优先使用正式素材包 02_ASSETS 及获接受的派生修复，禁止从本图抠取正式素材。'
                if other_product and node.text and node.text.startswith('仅用于理解历史视觉语言。'):
                    node.text = '本页为原交接文件中的其他产品图标，仅保留归档，不作为本游戏视觉、功能或素材依据；不得照搬或抠取。'
        body.insert(len(body) - 1, element)
    source_sha = hashlib.sha256(spec_path.read_bytes()).hexdigest()
    for style in doc.styles:
        if style.type in (1, 2):
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    for fonts in doc._element.findall('.//' + qn('w:rFonts')):
        fonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    doc.core_properties.subject = next(line for line in spec_path.read_text().splitlines()
                                       if line.startswith('版本：'))
    doc.core_properties.comments = 'SPEC source SHA-256: ' + source_sha
    doc.save(word_path)
    shutil.copyfile(word_path, second_path)
    return {'spec_sha256': source_sha, 'body_text_items': expected,
            'historical_images': len(doc.inline_shapes),
            'docx_sha256': hashlib.sha256(word_path.read_bytes()).hexdigest()}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--handoff', required=True, type=Path)
    parser.add_argument('--art-docs', required=True, type=Path)
    parser.add_argument('--audit', required=True, type=Path)
    args = parser.parse_args()
    name = '这也能叠_Codex开发交接_FINAL_v2_跨平台.docx'
    result = export(args.handoff / 'SPEC.md', args.handoff / name, args.art_docs / name)
    args.audit.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({key: value for key, value in result.items() if key != 'body_text_items'}, ensure_ascii=False))
